from pathlib import Path

import openpyxl

from amibroker_experiment_workbook import SHEET_NAMES, write_workbook


def _summary():
    passed = {
        "job_id": "valid-job",
        "strategy": "Demo",
        "periodicity": "Daily",
        "schedule": "native",
        "symbol": "ES",
        "sector": "Equity indexes",
        "median_profit_factor": 1.25,
        "median_car_mdd": .5,
        "html_path": "/tmp/detail.html",
        "json_path": "/tmp/detail.json",
    }
    return {
        "experiment_id": "experiment-1",
        "job_counts": {"total": 2, "admitted": 1, "invalid": 1},
        "passed_candidates": [passed],
        "low_touch_recommendations": [],
        "frequent_entry_exceptions": [],
        "all_symbol_results": [passed],
        "failed_or_incomplete": [{"job_id": "tampered-job", "reason": "hash mismatch"}],
        "schedule_comparisons": [],
    }


def test_workbook_has_exact_sheets_and_excludes_invalid_jobs(tmp_path):
    path = write_workbook(_summary(), tmp_path / "review.xlsx")
    book = openpyxl.load_workbook(path, data_only=False)
    assert book.sheetnames == SHEET_NAMES
    assert "tampered-job" not in repr(list(book["Passed Candidates"].values))
    assert "tampered-job" in repr(list(book["Failed or Incomplete"].values))
    assert book["Passed Candidates"].freeze_panes == "A2"
    assert book["Passed Candidates"].auto_filter.ref


def test_workbook_detail_paths_are_clickable(tmp_path):
    book_path = write_workbook(_summary(), tmp_path / "review.xlsx")
    book = openpyxl.load_workbook(book_path, data_only=False)
    sheet = book["Passed Candidates"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    html = sheet.cell(2, columns["html_path"])
    assert html.value == "/tmp/detail.html"
    assert html.hyperlink.target == Path("/tmp/detail.html").resolve().as_uri()
