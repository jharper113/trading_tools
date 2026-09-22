"""Create the compact review workbook for an AmiBroker experiment."""

from __future__ import annotations

import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


SHEET_NAMES = [
    "Passed Candidates",
    "Low-Touch Recommendations",
    "Frequent-Entry Exceptions",
    "All Symbol Results",
    "Failed or Incomplete",
    "Experiment Summary",
]

SYMBOL_COLUMNS = [
    "job_id", "strategy", "periodicity", "schedule", "symbol", "sector",
    "profitable_paramsets_pct", "median_profit_factor", "median_trades",
    "median_car_mdd", "median_max_drawdown_pct", "individual_pass",
    "sector_pass", "selected_representative", "html_path", "json_path",
]
COMPARISON_COLUMNS = [
    "strategy", "symbol", "native_job", "low_touch_job", "shared_grid_rows",
    "grid_coverage", "native_median_car_mdd", "low_touch_median_car_mdd",
    "relative_car_mdd_improvement", "profit_factor_improvement",
    "drawdown_reduction", "recommendation", "reason",
]
FAILURE_COLUMNS = ["job_id", "reason"]

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
YELLOW_FILL = PatternFill("solid", fgColor="FFF2CC")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")


def _cell_value(value):
    if isinstance(value, (list, dict, tuple, set)):
        return str(value)
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def prepare_sheet(book: Workbook, title: str, columns: list[str], rows: list[dict]):
    sheet = book.create_sheet(title)
    sheet.append(columns)
    for row in rows:
        sheet.append([_cell_value(row.get(column)) for column in columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    for column in sheet.columns:
        letter = column[0].column_letter
        width = min(55, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[letter].width = width
    for row in sheet.iter_rows(min_row=2):
        values = {columns[index]: cell.value for index, cell in enumerate(row)}
        decision = str(values.get("recommendation", ""))
        fill = None
        if values.get("selected_representative") is True or decision == "FREQUENT_ENTRY_WFA_CANDIDATE":
            fill = GREEN_FILL
        elif decision == "LOW_TOUCH":
            fill = YELLOW_FILL
        elif title == "Failed or Incomplete" or decision == "NOT COMPARABLE":
            fill = RED_FILL
        if fill:
            for cell in row:
                cell.fill = fill
        for key in ("html_path", "json_path"):
            if key in columns:
                cell = row[columns.index(key)]
                if cell.value:
                    cell.hyperlink = Path(str(cell.value)).resolve().as_uri()
                    cell.style = "Hyperlink"
        for index, name in enumerate(columns):
            if any(token in name for token in ("pct", "car_mdd", "improvement", "reduction", "coverage")):
                row[index].number_format = "0.0000"
    return sheet


def write_workbook(summary: dict, destination: Path) -> Path:
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.remove(book.active)
    prepare_sheet(book, SHEET_NAMES[0], SYMBOL_COLUMNS, summary.get("passed_candidates", []))
    prepare_sheet(book, SHEET_NAMES[1], COMPARISON_COLUMNS, summary.get("low_touch_recommendations", []))
    prepare_sheet(book, SHEET_NAMES[2], COMPARISON_COLUMNS, summary.get("frequent_entry_exceptions", []))
    prepare_sheet(book, SHEET_NAMES[3], SYMBOL_COLUMNS, summary.get("all_symbol_results", []))
    prepare_sheet(book, SHEET_NAMES[4], FAILURE_COLUMNS, summary.get("failed_or_incomplete", []))
    counts = summary.get("job_counts", {})
    summary_rows = [
        {"item": "experiment_id", "value": summary.get("experiment_id")},
        *({"item": f"jobs_{key}", "value": value} for key, value in counts.items()),
        {"item": "passed_candidates", "value": len(summary.get("passed_candidates", []))},
        {"item": "low_touch_recommendations", "value": len(summary.get("low_touch_recommendations", []))},
        {"item": "frequent_entry_exceptions", "value": len(summary.get("frequent_entry_exceptions", []))},
    ]
    prepare_sheet(book, SHEET_NAMES[5], ["item", "value"], summary_rows)
    book.save(destination)
    return destination
